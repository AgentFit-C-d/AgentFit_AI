"""Explicitly approved, redacted real-document Profile evaluation."""

import argparse
from dataclasses import dataclass
import hashlib
import json
from pathlib import Path
import re

from .diagnostics import safe_code
from .docling_structured_trial import convert_structured_pdf_bytes
from .false_complete_evaluation import write_safe_json
from .langextract_solar_trial import load_key
from .profile import FIELDS
from .solar import AnalysisError, SolarAnalyzer, post_solar_inline


ALLOWED_HASHES = {
    "H01": "57f13a242906bb5037b7fc65c5dbc30e42687436b27f9349699198449fc31239",
    "H02": "2a290fff891ac850a03ea715629a8d9fef1fb1628776eec8e852c8175b2bb24d",
    "H03": "a6ef17b02d392865c8251e0adb81c6dadab5781e57e53c418a4ddcf609b1136f",
}
KINDS = {"H01": "MARKDOWN", "H02": "DOCX", "H03": "PDF"}
_PERSONAL_DATA = (
    re.compile(r"(?<!\d)01[016789][\s-]?\d{3,4}[\s-]?\d{4}(?!\d)"),
    re.compile(r"(?<!\d)\d{8}(?!\d)"),
    re.compile(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}"),
)
_SAFE_STAGES = frozenset(("core", "features", "repair", "semantic_review",
                          "semantic_repair", "source_repair",
                          "section_features", "feature_curation",
                          "section_feature_review"))
_SAFE_OUTCOMES = frozenset(("started", "failed", "response_received",
                            "validated", "validation_failed", "semantic_failed"))


@dataclass(frozen=True)
class PreparedCase:
    id: str
    text: str
    checks: list[dict]
    source_sha256: str
    redacted_sha256: str
    page_count: int | None
    heading_count: int
    table_count: int


def redact_text(kind: str, text: str) -> str:
    if type(text) is not str or not text:
        raise ValueError("invalid source text")
    if kind == "MARKDOWN":
        first = "## 5. 팀 구성원 역할분배"
        next_sections = list(re.finditer(r"(?m)^## 6\.", text))
        if text.count(first) != 1 or len(next_sections) != 1:
            raise ValueError("markdown redaction boundary missing")
        start, end = text.index(first), next_sections[0].start()
        if not start < end:
            raise ValueError("markdown redaction boundary invalid")
        text = text[:start] + text[end:]
    elif kind == "PDF":
        marker = "프로젝트 메인 주제"
        if text.count(marker) != 1:
            raise ValueError("pdf redaction boundary missing")
        text = text[text.index(marker):]
    elif kind != "DOCX":
        raise ValueError("invalid source kind")
    text = text.strip()
    if (not text or len(text) > 100_000 or
            any(pattern.search(text) for pattern in _PERSONAL_DATA)):
        raise ValueError("private data or invalid redaction")
    return text


def _checks(value) -> list[dict]:
    if type(value) is not list or not value:
        raise ValueError("missing gold checks")
    ids = set()
    for check in value:
        if (type(check) is not dict or
                set(check) not in ({"id", "field", "expect_null"},
                                   {"id", "field", "contains_any"}) or
                type(check["id"]) is not str or
                re.fullmatch(r"C\d{2}", check["id"]) is None or
                check["id"] in ids or check["field"] not in FIELDS):
            raise ValueError("invalid gold check")
        ids.add(check["id"])
        if "expect_null" in check:
            if check["expect_null"] is not True:
                raise ValueError("invalid null check")
        elif (type(check["contains_any"]) is not list or
              not check["contains_any"] or
              any(type(term) is not str or not term.strip() or len(term) > 200
                  for term in check["contains_any"])):
            raise ValueError("invalid value check")
    return value


def prepare_cases(cases: list[dict], *, allowed_hashes=ALLOWED_HASHES,
                  pdf_converter=convert_structured_pdf_bytes) -> list[PreparedCase]:
    """Verify all approved source hashes and scrub all text before key loading."""
    if (type(cases) is not list or type(allowed_hashes) is not dict or
            len(cases) != len(allowed_hashes) or
            {case.get("id") for case in cases if type(case) is dict} !=
            set(allowed_hashes)):
        raise ValueError("invalid approved manifest")
    prepared = []
    for case in cases:
        if (type(case) is not dict or
                set(case) != {"id", "kind", "path", "sha256", "checks"} or
                case["sha256"] != allowed_hashes[case["id"]] or
                (case["id"] in KINDS and case["kind"] != KINDS[case["id"]]) or
                case["kind"] not in ("MARKDOWN", "DOCX", "PDF") or
                type(case["path"]) is not str or not case["path"]):
            raise ValueError("invalid approved source")
        checks = _checks(case["checks"])
        raw = Path(case["path"]).read_bytes()
        source_hash = hashlib.sha256(raw).hexdigest()
        if source_hash != case["sha256"]:
            raise ValueError("source hash mismatch")
        if case["kind"] == "MARKDOWN":
            text = raw.decode("utf-8-sig")
            page_count = None
            headings = sum(line.lstrip().startswith("#")
                           for line in text.splitlines())
            tables = 0
        elif case["kind"] == "DOCX":
            from docx import Document

            source = Document(Path(case["path"]))
            if source.tables or source.inline_shapes:
                raise ValueError("unsupported docx structure")
            text = "\n".join(paragraph.text for paragraph in source.paragraphs
                             if paragraph.text.strip())
            page_count = None
            headings = sum(paragraph.style.name.startswith("Heading")
                           for paragraph in source.paragraphs)
            tables = 0
        else:
            converted = pdf_converter(raw)
            text = converted.extracted.text
            page_count = converted.extracted.page_count
            headings = converted.heading_count
            tables = converted.table_count
        redacted = redact_text(case["kind"], text)
        prepared.append(PreparedCase(
            case["id"], redacted, checks, source_hash,
            hashlib.sha256(redacted.encode("utf-8")).hexdigest(),
            page_count, headings, tables))
    return prepared


def score_profile(profile: dict, checks: list[dict]) -> dict:
    if (type(profile) is not dict or type(profile.get("data")) is not dict):
        raise ValueError("invalid profile")
    checks = _checks(checks)
    failures = []
    for check in checks:
        value = profile["data"].get(check["field"])
        if "expect_null" in check:
            passed = value is None
        else:
            values = value if type(value) is list else [value]
            passed = any(type(item) is str and
                         any(term.casefold() in item.casefold()
                             for term in check["contains_any"])
                         for item in values)
        if not passed:
            failures.append(check["id"])
    return {"matched": len(checks)-len(failures), "total": len(checks),
            "failed_check_ids": failures}


def make_analyzer(key: str, timeout_seconds: int) -> SolarAnalyzer:
    if timeout_seconds not in (40, 120, 600):
        raise ValueError("unsupported diagnostic timeout")
    return SolarAnalyzer(
        key, transport=post_solar_inline,
        analysis_timeout_seconds=timeout_seconds,
        field_call_timeout_seconds=120 if timeout_seconds == 600 else 40,
        experimental_long_timeout=timeout_seconds == 600)


def safe_trace(diagnostic: dict | None) -> dict:
    if type(diagnostic) is not dict:
        return {"calls": []}
    trace = {"calls": []}
    if type(diagnostic.get("elapsed_ms")) is int:
        trace["elapsed_ms"] = diagnostic["elapsed_ms"]
    for call in diagnostic.get("calls", []):
        stage = call.get("stage")
        outcome = call.get("outcome")
        row = {"stage": stage if stage in _SAFE_STAGES else "other",
               "outcome": outcome if outcome in _SAFE_OUTCOMES else "other"}
        if "error" in call:
            row["error"] = safe_code(call["error"])
        trace["calls"].append(row)
    return trace


def select_cases(cases: list[PreparedCase], case_id: str | None) -> list[PreparedCase]:
    if case_id is None:
        return cases
    selected = [case for case in cases if case.id == case_id]
    if len(selected) != 1:
        raise ValueError("invalid case selection")
    return selected


def evaluate_cases(cases: list[PreparedCase], key: str,
                   *, analyzer_factory=None, timeout_seconds=40) -> dict:
    rows = []
    for case in cases:
        row = {"case_id": case.id, "source_sha256": case.source_sha256,
               "redacted_sha256": case.redacted_sha256,
               "characters": len(case.text), "page_count": case.page_count,
               "heading_count": case.heading_count,
               "table_count": case.table_count}
        try:
            analyzer = (analyzer_factory(key) if analyzer_factory else
                        make_analyzer(key, timeout_seconds))
            result = analyzer.analyze(case.text, case.id)
            row.update(score_profile(result.profile, case.checks))
            row.update(outcome="complete", provider_calls=result.provider_calls)
        except AnalysisError as error:
            row.update(outcome="failed", error=safe_code(error.code),
                       diagnostic=safe_trace(error.diagnostics))
        except Exception:
            row.update(outcome="failed", error="TRIAL_FAILURE")
        rows.append(row)
    complete = [row for row in rows if row["outcome"] == "complete"]
    return {"total": len(rows), "complete": len(complete),
            "failed": len(rows)-len(complete),
            "matched": sum(row["matched"] for row in complete),
            "checked": sum(row["total"] for row in complete),
            "rows": rows}


def main() -> int:
    parser = argparse.ArgumentParser(description="Redacted approved document holdout")
    parser.add_argument("--live", action="store_true")
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--manifest-sha256", required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--env-file", type=Path)
    parser.add_argument("--analysis-timeout-seconds", type=int,
                        choices=(40, 120, 600), default=40)
    parser.add_argument("--case-id", choices=tuple(KINDS))
    args = parser.parse_args()
    if not args.live:
        parser.error("--live required")
    if args.output.exists():
        parser.error("output already exists")
    try:
        manifest = args.manifest.read_bytes().replace(b"\r\n", b"\n")
        if (re.fullmatch(r"[0-9a-f]{64}", args.manifest_sha256) is None or
                hashlib.sha256(manifest).hexdigest() != args.manifest_sha256):
            raise ValueError("manifest hash mismatch")
        payload = json.loads(manifest.decode("utf-8"))
        if type(payload) is not dict or set(payload) != {"cases"}:
            raise ValueError("invalid manifest")
        prepared = prepare_cases(payload["cases"])
    except Exception:
        parser.error("private document preflight failed")
    key = load_key(args.env_file)
    result = evaluate_cases(select_cases(prepared, args.case_id), key,
                            timeout_seconds=args.analysis_timeout_seconds)
    result["analysis_timeout_seconds"] = args.analysis_timeout_seconds
    forbidden = (key, *(case.text for case in prepared),
                 *(str(case.get("path", "")) for case in payload["cases"]))
    args.output.parent.mkdir(parents=True, exist_ok=True)
    write_safe_json(args.output, result, forbidden_strings=forbidden)
    print(json.dumps({name: result[name] for name in
                      ("total", "complete", "failed", "matched", "checked")}))
    return 0 if result["failed"] == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
